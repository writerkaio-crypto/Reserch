"""中古品の画像差し替え候補 (仕入れ先と同じ / 近い状態のメルカリ出品) を Excel にまとめる。

  python make_condition_excel.py <判定JSON>

判定 JSON (人が選んだ候補) の形式:
  {"<SKU>": {"supplier_url": ..., "supplier_condition": ..., "supplier_note": ...,
             "note": SKU 全体の注意点,
             "picks": [{"url": ..., "grade": "◎" or "○", "note": ...}, ...]}}
候補の状態・画像枚数・価格などは output/mercari_condition/<SKU>.json (mercari_condition_search.py の結果) から引く。
出力: output/中古画像差し替え候補_YYYY-MM-DD.xlsx
"""

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

from mercari_condition_search import rank

OUT_DIR = Path(__file__).parent / "output"
DATA_DIR = OUT_DIR / "mercari_condition"
STATUS = {"on_sale": "販売中", "sold_out": "売り切れ", "trading": "取引中(売約済)",
          "shops": "メルカリShops"}
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
GRADE_FILL = {"◎": PatternFill("solid", fgColor="E2EFDA"), "○": PatternFill("solid", fgColor="FFF2CC")}
NONE_FILL = PatternFill("solid", fgColor="FCE4D6")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")


def load_items(sku):
    items = {}
    for path in DATA_DIR.glob(f"{sku}*.json"):
        for r in json.loads(path.read_text()):
            items[r["url"]] = r
    return items


def link(cell, url, text=None):
    cell.value = text or url
    cell.hyperlink = url
    cell.font = Font(color="0563C1", underline="single")


def write_header(ws, headers, widths):
    ws.append(headers)
    for i, (c, w) in enumerate(zip(ws[1], widths), 1):
        c.fill, c.font, c.border = HEAD_FILL, Font(bold=True), BORDER
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[c.column_letter].width = w
    ws.freeze_panes = "A2"


def main():
    review = json.loads(Path(sys.argv[1]).read_text())
    ebay = json.loads((DATA_DIR / "ebay_items.json").read_text())
    wb = Workbook()

    # --- SKU ごとのまとめ
    ws = wb.active
    ws.title = "SKUまとめ"
    write_header(ws, ["SKU", "eBay商品名", "eBay状態", "仕入先URL (AM〜AQ列の左端)", "仕入先の状態",
                      "候補数 ◎/○", "注意点"], [7, 45, 12, 40, 26, 11, 60])
    for sku, rv in review.items():
        e = ebay[sku]
        g = [p["grade"] for p in rv["picks"]]
        ws.append([sku, e["title"], e["condition"], None, rv["supplier_condition"],
                   f"{g.count('◎')} / {g.count('○')}" if g else "候補なし", rv.get("note", "")])
        row = ws[ws.max_row]
        link(row[3], rv["supplier_url"])
        for c in row:
            c.alignment, c.border = WRAP, BORDER
        if not g:
            for c in row:
                c.fill = NONE_FILL

    # --- 候補一覧 (1 行 = 1 候補)
    ws = wb.create_sheet("候補一覧")
    headers = ["反映\n(OK/NG)", "SKU", "一致度", "メルカリURL", "メルカリ商品名", "状態",
               "仕入先との差", "画像枚数", "販売状況", "価格(円)", "確認ポイント", "説明文(抜粋)",
               "仕入先の状態", "eBay出品(US)"]
    write_header(ws, headers, [9, 7, 7, 40, 40, 18, 10, 8, 13, 9, 45, 60, 22, 30])
    dv = DataValidation(type="list", formula1='"OK,NG"', allow_blank=True)
    ws.add_data_validation(dv)
    for sku, rv in review.items():
        items = load_items(sku)
        base = rank(rv["supplier_condition"].split("（")[0].split("(")[0])
        e = ebay[sku]
        ebay_url = f"https://www.ebay.com/itm/{e['item_id']}"
        if not rv["picks"]:
            ws.append([None, sku, "－", "候補なし", None, None, None, None, None, None,
                       rv.get("note", ""), None, rv["supplier_condition"], None])
            row = ws[ws.max_row]
            link(row[13], ebay_url)
            for c in row:
                c.alignment, c.border, c.fill = WRAP, BORDER, NONE_FILL
            continue
        for p in rv["picks"]:
            it = items[p["url"]]
            rk = rank(it["condition"])
            diff = "同じ" if base is None or rk == base else f"{rk - base:+d}段階"
            if base is None:
                diff = "（仕入先不明）"
            ws.append([None, sku, p["grade"], None, it["title"], it["condition"], diff,
                       len(it["images"]), STATUS.get(it["status"], it["status"]),
                       int(it["price"]) if str(it["price"]).isdigit() else it["price"],
                       p.get("note", ""), it["description"][:600], rv["supplier_condition"], None])
            row = ws[ws.max_row]
            link(row[3], p["url"])
            link(row[13], ebay_url)
            dv.add(row[0])
            for c in row:
                c.alignment, c.border = WRAP, BORDER
            row[2].fill = GRADE_FILL[p["grade"]]
            row[2].alignment = Alignment(horizontal="center", vertical="top")
    ws.auto_filter.ref = ws.dimensions

    # --- 判定基準
    ws = wb.create_sheet("判定基準")
    for line in [
        "■ 対象: 指定SKUのうち eBay の画像が2枚以下のもの（120 は eBay 側で既に10枚のため対象外）",
        "■ 仕入先: 販売管理シート AM〜AQ 列で一番左にある URL の商品の状態を基準にする",
        "■ メルカリ検索: 販売中・売り切れの両方。画像3枚以上の出品のみ",
        "■ 一致度",
        "   ◎ … 仕入先と同じ状態ランク、かつ同じ商品・同じ付属品構成",
        "   ○ … 状態ランクが1段階違うが、説明文の傷・使用感の内容が仕入先と近い",
        "   それ以外（2段階以上の差、別商品、付属品の過不足があるもの）は除外",
        "■ 状態ランク（メルカリ）: 新品、未使用 > 未使用に近い > 目立った傷や汚れなし > やや傷や汚れあり > 傷や汚れあり > 全体的に状態が悪い",
        "■ 注意: この環境からはメルカリの画像を表示できないため、判定は状態・商品名・説明文で行っています。",
        "   反映前に写真（付属品・傷の写り方）を目視で確認し、A列に OK / NG を記入してください。",
    ]:
        ws.append([line])
    ws.column_dimensions["A"].width = 120

    path = OUT_DIR / f"中古画像差し替え候補_{date.today().isoformat()}.xlsx"
    wb.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
