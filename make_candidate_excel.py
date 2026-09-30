"""画像差し替え候補 (image_candidates.json + image_review.json) を確認用 Excel にまとめる。

  python make_candidate_excel.py [SKU ...]

出力: output/画像差し替え候補_YYYY-MM-DD.xlsx
  「反映」列で OK を選んだ行だけを eBay に反映する想定。
"""

import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

OUT_DIR = Path(__file__).parent / "output"
COL_NAMES = {38: "AM", 39: "AN", 40: "AO", 41: "AP", 42: "AQ"}

FONT = "Meiryo UI"
F_BASE = Font(name=FONT, size=10)
F_HEAD = Font(name=FONT, size=10, bold=True, color="FFFFFF")
F_LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
F_TITLE = Font(name=FONT, size=13, bold=True)
F_NOTE = Font(name=FONT, size=9, color="666666")
FILL_HEAD = PatternFill("solid", fgColor="305496")
FILL_CHECK = PatternFill("solid", fgColor="FFF2CC")
FILL_OK = PatternFill("solid", fgColor="C6EFCE")
FILL_NG = PatternFill("solid", fgColor="F8CBAD")
FILL_NONE = PatternFill("solid", fgColor="EDEDED")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(vertical="top", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="top")


def selected_source(sku, cand, review):
    """(採用元, ページURL, 画像URL一覧) を返す。候補なしは None。"""
    pick = review.get("pick")
    if pick == "supplier":
        s = next(s for s in cand["suppliers"] if len(s.get("images") or []) >= 3)
        return f"仕入れ先（{COL_NAMES[s['col']]}列）", s["url"], s["images"]
    if pick:
        c = next(c for c in cand["sold"]["candidates"] if c["url"] == pick)
        return "メルカリ売り切れ（AS列）", c["url"], c["images"]
    return None


def put(ws, r, c, v, font=F_BASE, align=WRAP, link=None):
    cell = ws.cell(row=r, column=c, value=v)
    cell.font, cell.alignment, cell.border = font, align, BORDER
    if link:
        cell.hyperlink, cell.font = link, F_LINK
    return cell


def header(ws, names, widths, row=1):
    for c, (n, w) in enumerate(zip(names, widths), 1):
        cell = put(ws, row, c, n, F_HEAD, Alignment(horizontal="center", vertical="center",
                                                   wrap_text=True))
        cell.fill = FILL_HEAD
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.row_dimensions[row].height = 30


def main_sheet(ws, skus, cands, reviews):
    ws.title = "差し替え予定"
    max_imgs = max((len(sel[2]) for sku in skus
                    if (sel := selected_source(sku, cands[sku], reviews[sku]))), default=0)
    names = ["反映\n(OK/NG)", "SKU", "eBay商品番号", "eBayタイトル", "現在の\n画像枚数",
             "採用元", "採用ページ", "差し替え後\n画像枚数", "判定メモ"] + \
            [f"画像{i}" for i in range(1, max_imgs + 1)]
    header(ws, names, [9, 7, 15, 40, 9, 20, 34, 10, 42] + [9] * max_imgs)
    ws.freeze_panes = "E2"

    dv = DataValidation(type="list", formula1='"OK,NG"', allow_blank=True)
    ws.add_data_validation(dv)
    for r, sku in enumerate(skus, 2):
        c, rv = cands[sku], reviews[sku]
        sel = selected_source(sku, c, rv)
        check = put(ws, r, 1, None, align=CENTER)
        put(ws, r, 2, int(sku) if sku.isdigit() else sku, align=CENTER)
        put(ws, r, 3, c["item_id"], align=CENTER,
            link=f"https://www.ebay.com/itm/{c['item_id']}")
        put(ws, r, 4, c["ebay_title"])
        put(ws, r, 5, int(c["ebay_images"]), align=CENTER)
        if sel:
            source, url, images = sel
            check.fill = FILL_CHECK
            dv.add(check)
            put(ws, r, 6, source)
            put(ws, r, 7, url, link=url)
            put(ws, r, 8, len(images), align=CENTER)
            for i, u in enumerate(images, 10):
                put(ws, r, i, "開く", align=CENTER, link=u)
        else:
            check.value = "対象外"
            for col in range(1, 10):
                ws.cell(row=r, column=col).fill = FILL_NONE
            put(ws, r, 6, "候補なし").fill = FILL_NONE
            put(ws, r, 7, "").fill = FILL_NONE
            put(ws, r, 8, "").fill = FILL_NONE
        put(ws, r, 9, rv.get("memo", "")).fill = FILL_NONE if not sel else PatternFill()
        ws.row_dimensions[r].height = 45
    last = len(skus) + 1
    ws.conditional_formatting.add(f"A2:A{last}", CellIsRule(
        operator="equal", formula=['"OK"'], fill=FILL_OK))
    ws.conditional_formatting.add(f"A2:A{last}", CellIsRule(
        operator="equal", formula=['"NG"'], fill=FILL_NG))

    note = last + 2
    ws.cell(row=note, column=2, value="使い方").font = Font(name=FONT, size=10, bold=True)
    for i, t in enumerate([
        "・黄色の「反映」セルで OK / NG を選んでください。OK の行だけ eBay の画像を丸ごと差し替えます",
        "・「画像1〜」は差し替え後の画像（この順番で登録）。「開く」で画像を確認できます",
        "・灰色の行は条件に合う候補が見つからなかった SKU です（判定メモに理由）",
        "・候補の全件は「候補一覧」シートにあります",
    ], 1):
        ws.cell(row=note + i, column=2, value=t).font = F_BASE


def candidate_sheet(ws, skus, cands, reviews):
    ws.title = "候補一覧"
    header(ws, ["SKU", "種別", "ページ", "タイトル", "画像枚数", "採用", "備考"],
           [7, 20, 40, 50, 9, 7, 30])
    ws.freeze_panes = "A2"
    r = 2
    for sku in skus:
        c, rv = cands[sku], reviews[sku]
        sel = selected_source(sku, c, rv)
        chosen = sel[1] if sel else None
        rows = [(f"仕入れ先（{COL_NAMES[s['col']]}列）", s["url"], s.get("title", ""),
                 len(s.get("images") or []), s.get("error", "")) for s in c.get("suppliers", [])]
        sold = c.get("sold") or {}
        rows += [("メルカリ売り切れ", x["url"], x.get("title", ""), len(x["images"]), "")
                 for x in sold.get("candidates", [])]
        if sold:
            rows.append(("売り切れ検索", sold.get("search_url", ""),
                         f"売り切れ {sold.get('sold_found', 0)} 件中 {sold.get('checked', 0)} 件を確認",
                         None, sold.get("error", "")))
        for kind, url, title, n, err in rows:
            put(ws, r, 1, int(sku) if sku.isdigit() else sku, align=CENTER)
            put(ws, r, 2, kind)
            put(ws, r, 3, url, link=url or None)
            put(ws, r, 4, title)
            put(ws, r, 5, n, align=CENTER)
            put(ws, r, 6, "◎" if url == chosen else "", align=CENTER)
            put(ws, r, 7, err)
            if url == chosen:
                for col in range(1, 8):
                    ws.cell(row=r, column=col).fill = FILL_OK
            r += 1
    ws.auto_filter.ref = f"A1:G{r - 1}"


def main():
    cands = json.loads((OUT_DIR / "image_candidates.json").read_text())
    reviews = json.loads((OUT_DIR / "image_review.json").read_text())
    skus = sys.argv[1:] or sorted(reviews, key=lambda s: (len(s), s))
    wb = Workbook()
    main_sheet(wb.active, skus, cands, reviews)
    candidate_sheet(wb.create_sheet(), skus, cands, reviews)
    for ws in wb.worksheets:
        ws.page_setup.orientation = "landscape"
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
    path = OUT_DIR / f"画像差し替え候補_{date.today().isoformat()}.xlsx"
    wb.save(path)
    print(f"wrote {path}")


if __name__ == "__main__":
    main()
