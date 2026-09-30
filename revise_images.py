"""確認済み Excel (画像差し替え候補) で「反映」が OK の出品の画像だけを差し替える。

  python revise_images.py <確認済みExcel>            # 予行演習 (eBay は変更しない)
  python revise_images.py <確認済みExcel> --apply    # 実際に差し替える

環境変数 EBAY_USER_TOKEN (出品者本人のトークン) を使用。
  1) GetItem で出品中か・SKU が一致するかを確認し、現在の画像 URL をバックアップ
  2) UploadSiteHostedPictures で新しい画像を eBay の画像サーバーに取り込む
  3) ReviseFixedPriceItem で PictureDetails だけを送り、画像を丸ごと置き換える
     (タイトル・価格・説明など他の項目は送らないので変わらない)
結果: output/image_revise_log_YYYY-MM-DD.json (変更前の画像 URL を含む)
"""

import argparse
import json
import os
import sys
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

from openpyxl import load_workbook

from low_image_listings import NS, trading_call

OUT_DIR = Path(__file__).parent / "output"
FIRST_IMAGE_COL = 10  # 「画像1」列


def approved_rows(path):
    ws = load_workbook(path)["差し替え予定"]
    rows = []
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip().upper() != "OK":
            continue
        images = [ws.cell(r, c).hyperlink.target
                  for c in range(FIRST_IMAGE_COL, ws.max_column + 1)
                  if ws.cell(r, c).hyperlink]
        rows.append({"sku": str(ws.cell(r, 2).value), "item_id": str(ws.cell(r, 3).value),
                     "source": ws.cell(r, 7).value, "images": images})
    return rows


def get_item(token, item_id):
    root = trading_call(token, "GetItem", (
        f"<ItemID>{item_id}</ItemID>"
        "<OutputSelector>Item.ItemID</OutputSelector>"
        "<OutputSelector>Item.SKU</OutputSelector>"
        "<OutputSelector>Item.Site</OutputSelector>"
        "<OutputSelector>Item.SellingStatus.ListingStatus</OutputSelector>"
        "<OutputSelector>Item.PictureDetails</OutputSelector>"))
    it = root.find("e:Item", NS)
    return {
        "sku": it.findtext("e:SKU", "", namespaces=NS),
        "site": it.findtext("e:Site", "", namespaces=NS),
        "status": it.findtext("e:SellingStatus/e:ListingStatus", "", namespaces=NS),
        "pictures": [p.text for p in it.iterfind("e:PictureDetails/e:PictureURL", NS)],
    }


def upload_picture(token, url):
    root = trading_call(token, "UploadSiteHostedPictures", (
        f"<ExternalPictureURL>{escape(url)}</ExternalPictureURL>"
        "<PictureSet>Supersize</PictureSet>"))
    return root.findtext("e:SiteHostedPictureDetails/e:FullURL", namespaces=NS)


def revise_pictures(token, item_id, urls):
    pics = "".join(f"<PictureURL>{escape(u)}</PictureURL>" for u in urls)
    trading_call(token, "ReviseFixedPriceItem", (
        f"<Item><ItemID>{item_id}</ItemID>"
        f"<PictureDetails>{pics}</PictureDetails></Item>"))


def sku_matches(excel_sku, ebay_sku):
    a, b = excel_sku.strip(), ebay_sku.strip()
    return a == b or (a.isdigit() and b.isdigit() and int(a) == int(b))


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("excel")
    ap.add_argument("--apply", action="store_true", help="実際に eBay を変更する")
    args = ap.parse_args()
    token = os.environ.get("EBAY_USER_TOKEN", "")
    if len(token) < 100:
        sys.exit("EBAY_USER_TOKEN が未設定か途中で切れています")

    rows = approved_rows(args.excel)
    print(f"反映 OK: {len(rows)} 件 ({'本番' if args.apply else '予行演習'})")
    log_path = OUT_DIR / f"image_revise_log_{datetime.now():%Y-%m-%d}.json"
    log = json.loads(log_path.read_text()) if log_path.exists() else []

    for row in rows:
        entry = {"time": datetime.now().isoformat(timespec="seconds"), **row}
        try:
            cur = get_item(token, row["item_id"])
            entry["before"] = cur["pictures"]
            if cur["status"] != "Active":
                raise RuntimeError(f"出品中ではありません ({cur['status']})")
            if cur["site"] != "US":
                raise RuntimeError(f"US サイトの出品ではありません ({cur['site']})")
            if not sku_matches(row["sku"], cur["sku"]):
                raise RuntimeError(f"SKU が一致しません (eBay: {cur['sku']})")
            if not row["images"]:
                raise RuntimeError("画像 URL がありません")
            if args.apply:
                hosted = [upload_picture(token, u) for u in row["images"]]
                revise_pictures(token, row["item_id"], hosted)
                after = get_item(token, row["item_id"])["pictures"]
                entry.update(result="変更済み", after=after)
            else:
                entry["result"] = "予行演習OK"
            print(f"  SKU {row['sku']} ({row['item_id']}): {len(cur['pictures'])}枚 → "
                  f"{len(row['images'])}枚  {entry['result']}")
        except Exception as e:
            entry["result"] = f"エラー: {e}"
            print(f"  SKU {row['sku']} ({row['item_id']}): {entry['result']}")
        log.append(entry)
        if args.apply:
            log_path.write_text(json.dumps(log, ensure_ascii=False, indent=1))

    if args.apply:
        print(f"wrote {log_path}")


if __name__ == "__main__":
    main()
