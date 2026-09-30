"""画像差し替えシートの「OK」行について eBay 出品の画像を丸ごと差し替える。

使い方:
  python update_images.py <差し替えシート.xlsx>           # 事前チェックのみ (eBay は変更しない)
  python update_images.py <差し替えシート.xlsx> --apply   # 実際に差し替える

認証 (どちらか):
  EBAY_REFRESH_TOKEN + EBAY_CLIENT_ID / EBAY_CLIENT_SECRET  (推奨: 18か月有効)
  EBAY_USER_TOKEN  (OAuth ユーザートークン。2時間で失効)

処理:
  1. シート「差し替え予定」の A列=OK の行から商品番号と画像1〜15 のリンクを読む
  2. 画像をダウンロードして検証 (JPEG/PNG・長辺 500px 以上)
  3. --apply 時: GetItem で現在の画像を output/image_backup_<日付>.json に保存
     → UploadSiteHostedPictures で eBay に画像をアップロード
     → ReviseFixedPriceItem / ReviseItem で PictureDetails を差し替え
  4. 結果を Z列「反映結果」に書いたコピーを output/ に保存
"""

import argparse
import base64
import io
import json
import os
import re
import sys
import urllib.error
import urllib.parse
import urllib.request
import uuid
from datetime import datetime
from pathlib import Path
from xml.sax.saxutils import escape

import openpyxl
from PIL import Image

TRADING_URL = "https://api.ebay.com/ws/api.dll"
TOKEN_URL = "https://api.ebay.com/identity/v1/oauth2/token"
COMPAT = "1293"
SITE_ID = "0"  # US
SHEET = "差し替え予定"
IMG_COLS = range(11, 26)  # K..Y = 画像1..15
MIN_EDGE = 500
OUT_DIR = Path(__file__).parent / "output"
NS = "{urn:ebay:apis:eBLBaseComponents}"


def get_token():
    refresh = os.environ.get("EBAY_REFRESH_TOKEN")
    if refresh:
        cid = os.environ["EBAY_CLIENT_ID"]
        secret = os.environ["EBAY_CLIENT_SECRET"]
        scopes = os.environ.get(
            "EBAY_SCOPES",
            "https://api.ebay.com/oauth/api_scope "
            "https://api.ebay.com/oauth/api_scope/sell.inventory",
        )
        auth = base64.b64encode(f"{cid}:{secret}".encode()).decode()
        body = urllib.parse.urlencode({
            "grant_type": "refresh_token",
            "refresh_token": refresh,
            "scope": scopes,
        }).encode()
        req = urllib.request.Request(TOKEN_URL, data=body, headers={
            "Authorization": f"Basic {auth}",
            "Content-Type": "application/x-www-form-urlencoded",
        })
        with urllib.request.urlopen(req, timeout=30) as r:
            return json.load(r)["access_token"]
    return os.environ["EBAY_USER_TOKEN"]


def trading(token, call, xml_body, files=None):
    """Trading API を呼び出して (ack, errors, raw_xml) を返す。"""
    headers = {
        "X-EBAY-API-CALL-NAME": call,
        "X-EBAY-API-SITEID": SITE_ID,
        "X-EBAY-API-COMPATIBILITY-LEVEL": COMPAT,
        "X-EBAY-API-IAF-TOKEN": token,
    }
    xml = (
        '<?xml version="1.0" encoding="utf-8"?>'
        f'<{call}Request xmlns="urn:ebay:apis:eBLBaseComponents">'
        f"{xml_body}<ErrorLanguage>en_US</ErrorLanguage>"
        f"<WarningLevel>High</WarningLevel></{call}Request>"
    )
    if files:
        boundary = uuid.uuid4().hex
        parts = [
            f"--{boundary}\r\nContent-Disposition: form-data; "
            f'name="XML Payload"\r\nContent-Type: text/xml\r\n\r\n'.encode()
            + xml.encode() + b"\r\n"
        ]
        for name, data in files:
            parts.append(
                f"--{boundary}\r\nContent-Disposition: form-data; "
                f'name="image"; filename="{name}"\r\n'
                "Content-Transfer-Encoding: binary\r\n"
                "Content-Type: application/octet-stream\r\n\r\n".encode()
                + data + b"\r\n"
            )
        parts.append(f"--{boundary}--\r\n".encode())
        payload = b"".join(parts)
        headers["Content-Type"] = f"multipart/form-data; boundary={boundary}"
    else:
        payload = xml.encode()
        headers["Content-Type"] = "text/xml"
    req = urllib.request.Request(TRADING_URL, data=payload, headers=headers)
    with urllib.request.urlopen(req, timeout=120) as r:
        raw = r.read().decode()
    ack = (re.search(r"<Ack>(.*?)</Ack>", raw) or [None, ""])[1]
    errors = []
    for block in re.findall(r"<Errors>(.*?)</Errors>", raw, re.S):
        msg = re.search(r"<LongMessage>(.*?)</LongMessage>", block, re.S)
        sev = re.search(r"<SeverityCode>(.*?)</SeverityCode>", block)
        errors.append(f"[{sev[1] if sev else '?'}] {msg[1] if msg else block[:200]}")
    return ack, errors, raw


def read_targets(path):
    wb = openpyxl.load_workbook(path)
    ws = wb[SHEET]
    targets = []
    for r in range(2, ws.max_row + 1):
        if str(ws.cell(r, 1).value or "").strip().upper() != "OK":
            continue
        item_id = str(ws.cell(r, 3).value or "").strip()
        urls = [
            ws.cell(r, c).hyperlink.target
            for c in IMG_COLS if ws.cell(r, c).hyperlink
        ]
        targets.append({
            "row": r,
            "sku": str(ws.cell(r, 2).value or "").strip(),
            "item_id": item_id,
            "title": ws.cell(r, 4).value,
            "expected": ws.cell(r, 8).value,
            "urls": urls,
        })
    return wb, ws, targets


def download(url):
    req = urllib.request.Request(url, headers={"User-Agent": "Mozilla/5.0"})
    with urllib.request.urlopen(req, timeout=60) as r:
        data = r.read()
    img = Image.open(io.BytesIO(data))
    img.load()
    if img.format not in ("JPEG", "PNG"):
        raise ValueError(f"未対応の形式 {img.format}")
    if max(img.size) < MIN_EDGE:
        raise ValueError(f"解像度不足 {img.size}")
    return data, img.size


def get_item(token, item_id):
    ack, errors, raw = trading(
        token, "GetItem",
        f"<ItemID>{item_id}</ItemID><DetailLevel>ReturnAll</DetailLevel>",
    )
    if ack not in ("Success", "Warning"):
        raise RuntimeError("GetItem 失敗: " + "; ".join(errors))
    return {
        "listing_type": (re.search(r"<ListingType>(.*?)</ListingType>", raw) or [None, ""])[1],
        "status": (re.search(r"<ListingStatus>(.*?)</ListingStatus>", raw) or [None, ""])[1],
        "sku": (re.search(r"<Item>.*?<SKU>(.*?)</SKU>", raw, re.S) or [None, ""])[1],
        "pictures": re.findall(r"<PictureURL>(.*?)</PictureURL>", raw),
        "has_variations": "<Variations>" in raw,
    }


def upload_picture(token, name, data):
    ack, errors, raw = trading(
        token, "UploadSiteHostedPictures",
        f"<PictureName>{escape(name)}</PictureName>"
        "<PictureSet>Supersize</PictureSet>",
        files=[(name, data)],
    )
    url = re.search(r"<FullURL>(.*?)</FullURL>", raw)
    if ack not in ("Success", "Warning") or not url:
        raise RuntimeError("画像アップロード失敗: " + "; ".join(errors))
    return url[1]


def revise_pictures(token, item_id, listing_type, eps_urls):
    call = "ReviseFixedPriceItem" if listing_type == "FixedPriceItem" else "ReviseItem"
    pics = "".join(f"<PictureURL>{escape(u)}</PictureURL>" for u in eps_urls)
    ack, errors, _ = trading(
        token, call,
        f"<Item><ItemID>{item_id}</ItemID>"
        f"<PictureDetails>{pics}</PictureDetails></Item>",
    )
    if ack not in ("Success", "Warning"):
        raise RuntimeError(f"{call} 失敗: " + "; ".join(errors))
    return errors


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("xlsx")
    ap.add_argument("--apply", action="store_true", help="eBay を実際に更新する")
    ap.add_argument("--only", nargs="*", help="この SKU だけ処理する")
    args = ap.parse_args()

    wb, ws, targets = read_targets(args.xlsx)
    if args.only:
        targets = [t for t in targets if t["sku"] in args.only]
    print(f"OK 行: {len(targets)} 件  ({'本番' if args.apply else '事前チェック'})")

    token = get_token() if args.apply else None
    OUT_DIR.mkdir(exist_ok=True)
    stamp = datetime.now().strftime("%Y-%m-%d_%H%M")
    backup_path = OUT_DIR / f"image_backup_{stamp}.json"
    backups, results = [], {}

    for t in targets:
        label = f"SKU {t['sku']} / {t['item_id']}"
        try:
            if not t["urls"]:
                raise ValueError("画像リンクがありません")
            if t["expected"] and len(t["urls"]) != int(t["expected"]):
                raise ValueError(
                    f"画像数不一致 (シート {t['expected']} / リンク {len(t['urls'])})"
                )
            images = []
            for i, url in enumerate(t["urls"], 1):
                data, size = download(url)
                images.append((f"{t['sku']}_{i}.jpg", data))
            print(f"  {label}: 画像 {len(images)} 枚 OK")

            if not args.apply:
                results[t["row"]] = f"チェックOK ({len(images)}枚)"
                continue

            cur = get_item(token, t["item_id"])
            backups.append({"sku": t["sku"], "item_id": t["item_id"], **cur})
            backup_path.write_text(json.dumps(backups, ensure_ascii=False, indent=2))
            if cur["status"] != "Active":
                raise RuntimeError(f"出品が Active ではありません ({cur['status']})")
            if cur["sku"] and cur["sku"] != t["sku"]:
                raise RuntimeError(f"SKU 不一致 (eBay {cur['sku']} / シート {t['sku']})")

            eps = [upload_picture(token, n, d) for n, d in images]
            warns = revise_pictures(token, t["item_id"], cur["listing_type"], eps)
            msg = f"完了 {len(cur['pictures'])}→{len(eps)}枚"
            if warns:
                msg += " (警告: " + "; ".join(warns)[:200] + ")"
            results[t["row"]] = msg
            print(f"  {label}: {msg}")
        except (urllib.error.URLError, ValueError, RuntimeError, OSError) as e:
            results[t["row"]] = f"エラー: {e}"
            print(f"  {label}: エラー {e}", file=sys.stderr)

    ws.cell(1, 26, "反映結果")
    ws.cell(1, 26)._style = ws.cell(1, 25)._style
    for row, msg in results.items():
        ws.cell(row, 26, f"{msg} [{stamp}]")
    ws.column_dimensions["Z"].width = 40
    out = OUT_DIR / f"画像差し替え結果_{stamp}.xlsx"
    wb.save(out)
    print(f"wrote {out}")
    if backups:
        print(f"wrote {backup_path} (元画像URLのバックアップ)")
    if any(m.startswith("エラー") for m in results.values()):
        sys.exit(1)


if __name__ == "__main__":
    main()
